"""Structural checks for docs/deck/pitch_template.pptx, without rendering.

Run from the repo root after editing the deck by hand or rebuilding it:

    .venv/bin/python docs/deck/check_deck.py

It reports:
- text smaller than 14 pt, other than the source line and slide number;
- shapes that run off the slide or into the footer;
- text boxes whose estimated wrapped height exceeds the box;
- content shapes that overlap, other than a shape drawn inside another;
- main slides without a chart, table, picture, or drawn diagram, or with more than three bullets;
- notes without a script, or with [[ placeholders on slides.

The overflow estimate uses Arial's average character width, so treat it as a
prompt to look, not a verdict. The PNG previews are the final check.
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu, Inches, Pt

DECK = Path(__file__).resolve().parents[2] / "docs/deck/pitch_template.pptx"
MIN_PT = 14
SMALL_OK = {"Source line", "Slide number"}
FOOTER = {"Source line", "Slide number", "Rule"}
FOOTER_Y = Inches(6.86)  # the footer rule; content must end above it
CHAR_W = 0.50  # Arial average glyph width as a share of the font size; bold runs use 0.55
LINE_H = 1.2  # line height as a multiple of the font size


def est_height(shape):
    """Estimated wrapped text height in EMU."""
    tf = shape.text_frame
    width_pt = Emu(shape.width - tf.margin_left - tf.margin_right).pt
    total = 0.0
    for p in tf.paragraphs:
        runs = [r for r in p.runs if r.text]
        if not runs:
            continue
        size = max((r.font.size or Pt(18)).pt for r in runs)
        chars = sum(len(r.text) * (0.55 if r.font.bold else CHAR_W) for r in runs)
        indent = 0.32 * 72 if p._p.pPr is not None and p._p.pPr.get("marL") else 0
        per_line = max(1.0, (width_pt - indent) / size)
        lines = max(1, -(-chars // per_line))
        total += lines * size * LINE_H + (p.space_after.pt if p.space_after else 0) + (p.space_before.pt if p.space_before else 0)
    return Pt(total + Emu(tf.margin_top + tf.margin_bottom).pt)


def box(s):
    return s.left, s.top, s.left + s.width, s.top + s.height


def contains(a, b, slack=Emu(12700)):
    return a[0] - slack <= b[0] and a[1] - slack <= b[1] and a[2] + slack >= b[2] and a[3] + slack >= b[3]


def overlap(a, b):
    w = min(a[2], b[2]) - max(a[0], b[0])
    h = min(a[3], b[3]) - max(a[1], b[1])
    return max(0, w) * max(0, h)


def main():
    prs = Presentation(DECK)
    W, H = prs.slide_width, prs.slide_height
    problems = []
    for n, slide in enumerate(prs.slides, start=1):
        label = f"slide {n}"
        shapes = list(slide.shapes)
        visual = 0
        for s in shapes:
            if s.has_chart or s.shape_type == MSO_SHAPE_TYPE.PICTURE or s.has_table:
                visual += 1
            if s.left < 0 or s.top < 0 or s.left + s.width > W or s.top + s.height > H:
                problems.append(f"{label}: '{s.name}' runs off the slide")
            if s.name not in FOOTER and s.top + s.height > FOOTER_Y:
                problems.append(f"{label}: '{s.name}' reaches into the footer (bottom at "
                                f"{Emu(s.top + s.height).inches:.2f} in, rule at {Emu(FOOTER_Y).inches:.2f} in)")
            if s.has_text_frame and s.text_frame.text.strip():
                if "[[" in s.text_frame.text:
                    problems.append(f"{label}: placeholder on slide: {s.text_frame.text.strip()[:70]}")
                for p in s.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size and r.font.size.pt < MIN_PT and s.name not in SMALL_OK:
                            problems.append(f"{label}: '{s.name}' has {r.font.size.pt:.0f} pt text: {r.text[:40]}")
                if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX or s.name in ("Placeholder",):
                    need = est_height(s)
                    if need > s.height * 1.08:
                        problems.append(f"{label}: '{s.name}' may overflow (needs about {Emu(need).inches:.2f} in, "
                                        f"has {Emu(s.height).inches:.2f} in): {s.text_frame.text.strip()[:50]}")
            if s.name == "Bullets" and len([p for p in s.text_frame.paragraphs if p.text.strip()]) > 3:
                problems.append(f"{label}: more than three bullets")
        # Overlaps between content shapes: charts, tables, pictures, and text boxes.
        content = [s for s in shapes if s.has_chart or s.has_table or s.shape_type in
                   (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.TEXT_BOX)]
        for i, a in enumerate(content):
            for b in content[i + 1:]:
                ba, bb = box(a), box(b)
                if contains(ba, bb) or contains(bb, ba):
                    continue
                area = overlap(ba, bb)
                if area > 0.02 * min(a.width * a.height, b.width * b.height):
                    problems.append(f"{label}: '{a.name}' overlaps '{b.name}'")
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        if len(notes.split("\n\n")[0].split()) < 40:
            problems.append(f"{label}: the notes don't open with a speaker script of 40 or more words")
        drawn = sum(1 for s in shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE)
        if n <= 9 and visual == 0 and drawn < 5:  # five or more drawn shapes count as a native diagram
            problems.append(f"{label}: main slide without a chart, table, picture, or diagram")
    print(f"{DECK.name}: {len(prs.slides)} slides")
    for line in problems:
        print("  " + line)
    print("no problems found" if not problems else f"{len(problems)} possible problems")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
