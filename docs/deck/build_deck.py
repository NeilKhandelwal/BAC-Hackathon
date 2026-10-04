"""Build the editable pitch template: docs/deck/pitch_template.pptx.

Run from the repo root:

    .venv/bin/python docs/deck/build_deck.py

Reads docs/figures/facts.json, scratch/weighting/out/smaa_acceptability.csv,
and PNGs in docs/figures/. Writes the deck, docs/deck/numbers_to_check.md,
and docs/deck/sources.md. Numbers that come from facts.json refresh on a
rebuild. Numbers from research write-ups are typed in below, next to the file
they come from.

A rebuild overwrites the .pptx. Make wording changes here, or edit the .pptx
by hand once the numbers are final.
"""
import csv
import json
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_AXIS_CROSSES, XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
DECK = ROOT / "docs/deck"
FIG = ROOT / "docs/figures"
FACTS = json.loads((FIG / "facts.json").read_text())

# ---------------------------------------------------------------------------
# Theme: the cockpit palette from DESIGN.md, so slides and demo match.
# ---------------------------------------------------------------------------
FONT = "Arial"
INK = RGBColor(0x11, 0x15, 0x18)
INK2 = RGBColor(0x46, 0x50, 0x5A)
INK3 = RGBColor(0x5F, 0x6A, 0x74)
RULE = RGBColor(0xD3, 0xD8, 0xD6)
GROUND = RGBColor(0xF3, 0xF5, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY = RGBColor(0x0B, 0x3C, 0x6F)
BLUE = RGBColor(0x4F, 0x82, 0xBD)
PALE = RGBColor(0xDC, 0xE3, 0xEA)
GREY = RGBColor(0x8E, 0x8F, 0x8A)
EMBER = RGBColor(0xD9, 0x53, 0x2B)
AMBER = RGBColor(0xB8, 0x62, 0x00)
GREEN = RGBColor(0x23, 0x80, 0x4F)

PILLARS = [  # key, label, hue
    ("energy_carbon", "Energy & carbon", RGBColor(0xB8, 0x62, 0x00)),
    ("water", "Water", RGBColor(0x14, 0x75, 0x8F)),
    ("climate_resilience", "Climate", RGBColor(0xB0, 0x30, 0x3F)),
    ("grid_infrastructure", "Grid", RGBColor(0x5B, 0x4B, 0xA0)),
    ("land", "Land", RGBColor(0x5E, 0x6E, 0x12)),
    ("community", "Community", RGBColor(0x23, 0x80, 0x4F)),
    ("permitting", "Permitting", RGBColor(0x86, 0x52, 0x3A)),
    ("cost", "Cost of power", RGBColor(0x9A, 0x3A, 0x8C)),
]

W, H = Inches(13.333), Inches(7.5)
MX = Inches(0.6)  # side margin
CW = W - 2 * MX  # content width
TOP = Inches(2.2)  # content top, below a two-line headline
BOTTOM = Inches(6.72)  # content bottom

# ---------------------------------------------------------------------------
# Registries: numbers to check, sources, placeholders.
# ---------------------------------------------------------------------------
CHECKS = []  # (slide, what, value, file, field)
SOURCES = {}  # slide -> list of source lines for sources.md
PLACEHOLDERS = []  # (slide, text)
SCRIPT_WORDS = {}  # slide -> words in the spoken script


def fmt(n, d=0):
    return f"{n:,.{d}f}"


def pct(x, d=1):
    return f"{x * 100:.{d}f}%"


# ---------------------------------------------------------------------------
# Quotes. Each one is verified at the URL given, or it becomes a placeholder.
# ---------------------------------------------------------------------------
# Verified 2026-10-04 by opening each primary source and matching the words exactly (ignoring whitespace and
# straight versus curly quotes). An ellipsis marks an omission. Unused quotes stay in the bank for sources.md.
QUOTE_BANK = [
    dict(key="ny_eo62", slide="1",
         text="New Yorkers have expressed legitimate concerns regarding the potential impacts of the siting and "
              "operation of data centers on energy use, water use, water quality…",
         who="New York Executive Order No. 62, Gov. Kathy Hochul, July 14, 2026", cite="New York Executive Order 62, July 2026",
         source="Executive Order No. 62, WHEREAS clause 9",
         url="https://www.governor.ny.gov/executive-order/no-62-establishing-temporary-moratorium-data-centers-new-york-while-state-develops"),
    dict(key="wa_water", slide="4",
         text="The direct water requirements of data centers can be substantial, depending on the size and type of "
              "cooling system used.",
         who="Washington Data Center Workgroup, Preliminary Report, Dec. 2025", cite="Washington Data Center Workgroup, Dec. 2025",
         source="Data Center Workgroup: Preliminary Report (Executive Order 25-05), Finding 19, p. 13",
         url="https://dor.wa.gov/sites/default/files/2025-12/2025DataCntrWrkgrpPrelimReport.pdf"),
    dict(key="pud_speed", slide="7",
         text="…service cannot always be provided as quickly as customers or developers may prefer.",
         who="Grant County PUD, Data Center FAQs, Aug. 28, 2026", cite="Grant County PUD, Aug. 2026",
         source="Grant PUD & Data Centers: FAQs, Q6",
         url="https://www.grantpud.org/blog/data-center-faqs"),
    dict(key="pud_pays", slide="8",
         text="Large power users, including data centers, are responsible for paying the costs of additional "
              "generation, transmission, and distribution infrastructure needed specifically to serve their loads.",
         who="Grant County PUD, Data Center FAQs, Aug. 28, 2026", cite="Grant County PUD, Aug. 2026",
         source="Grant PUD & Data Centers: FAQs, Q5",
         url="https://www.grantpud.org/blog/data-center-faqs"),
    # Verified, not on a slide. Use in speaker notes, Q&A, or a swap.
    dict(key="iea_us_growth", slide="unused",
         text="In the United States, data centres account for nearly half of electricity demand growth between now and 2030.",
         who="International Energy Agency, Energy and AI, April 10, 2025 (CC BY 4.0)",
         source="Energy and AI, Executive summary", url="https://www.iea.org/reports/energy-and-ai/executive-summary"),
    dict(key="doe_growth", slide="unused",
         text="The report estimates that data center load growth has tripled over the past decade and is projected to "
              "double or triple by 2028.",
         who="U.S. Department of Energy, Dec. 20, 2024",
         source="DOE Releases New Report Evaluating Increase in Electricity Demand from Data Centers, paragraph 1",
         url="https://www.energy.gov/articles/doe-releases-new-report-evaluating-increase-electricity-demand-data-centers"),
    dict(key="lbnl_2023", slide="1 (chart values)",
         text="U.S. data center energy use has continued to grow at an increasing rate, reaching 176 TWh by 2023, "
              "representing 4.4% of total U.S. electricity consumption.",
         who="Shehabi et al., Lawrence Berkeley National Laboratory, Dec. 2024",
         source="2024 United States Data Center Energy Usage Report (LBNL-2001637), Executive Summary, p. 5",
         url="https://eta-publications.lbl.gov/sites/default/files/2024-12/lbnl-2024-united-states-data-center-energy-usage-report_1.pdf"),
    dict(key="lbnl_2028", slide="1 (chart values)",
         text="This annual energy use also represents 6.7% to 12.0% of total U.S. electricity consumption forecasted for 2028.",
         who="Shehabi et al., Lawrence Berkeley National Laboratory, Dec. 2024",
         source="2024 United States Data Center Energy Usage Report, Executive Summary, p. 6",
         url="https://eta-publications.lbl.gov/sites/default/files/2024-12/lbnl-2024-united-states-data-center-energy-usage-report_1.pdf"),
    dict(key="lbnl_water", slide="unused",
         text="However, the water evaporation process in cooling towers has raised concerns regarding data center water "
              "consumption and availability at the local level.",
         who="Shehabi et al., Lawrence Berkeley National Laboratory, Dec. 2024",
         source="2024 United States Data Center Energy Usage Report, cooling-system table, p. 41",
         url="https://eta-publications.lbl.gov/sites/default/files/2024-12/lbnl-2024-united-states-data-center-energy-usage-report_1.pdf"),
    dict(key="wa_load_growth", slide="unused (slide 7 notes)",
         text="Data centers are the largest source of expected load growth in the Pacific Northwest.",
         who="Washington Data Center Workgroup, Preliminary Report, Dec. 2025",
         source="Data Center Workgroup: Preliminary Report, Finding 6, p. 10",
         url="https://dor.wa.gov/sites/default/files/2025-12/2025DataCntrWrkgrpPrelimReport.pdf"),
    dict(key="wa_utc", slide="unused",
         text="Planning now will help the UTC protect customers and ensure electric services stay safe, equitable, "
              "available, reliable, and fairly priced as demand grows.",
         who="Washington Utilities and Transportation Commission, July 28, 2026 (docket UE-260162; covers "
             "investor-owned utilities, not Grant PUD)",
         source="Media advisory: UTC workshop on emerging large electric loads, paragraph 2",
         url="https://www.utc.wa.gov/news/2026/media-advisory-public-invited-join-utc-workshop-emerging-large-electric-loads"),
    dict(key="bpa_nlsl", slide="unused (slide 5 notes)",
         text="…the customer must serve that load, and any increases to it, with either power from Bonneville at "
              "the NR rate or a dedicated nonfederal resource.",
         who="Bonneville Power Administration, Fact Sheet: New large single load, Oct. 2020",
         source="Fact sheet DOE/BP-5045, p. 2",
         url="https://www.bpa.gov/-/media/Aep/about/publications/fact-sheets/fs-202011-New-Large-Single-Load.pdf"),
    dict(key="wa_siting", slide="unused (slide 9 notes, alternative close)",
         text="Developers can avoid and minimize environmental and other community impacts through coordinated "
              "planning … when designing projects and choosing project sites.",
         who="Washington Data Center Workgroup, Preliminary Report, Dec. 2025",
         source="Data Center Workgroup: Preliminary Report, Finding 18b, p. 13",
         url="https://dor.wa.gov/sites/default/files/2025-12/2025DataCntrWrkgrpPrelimReport.pdf"),
]
QUOTES = {q["key"]: q for q in QUOTE_BANK}

# Openly licensed photos, licenses read on each Commons file page on 2026-10-04. Not placed on a slide:
# every main slide already carries a chart. Listed in sources.md for a swap.
PHOTOS = {
    "Wanapum Dam": dict(credit="Photo: Williamborg, public domain, via Wikimedia Commons", license="Public domain",
                        page="https://commons.wikimedia.org/wiki/File:Wanapum_Dam_from_West_Shore_-_downstream_10360031.jpg"),
    "Quincy, WA aerial": dict(credit="Photo: Tedder, CC BY-SA 4.0, via Wikimedia Commons", license="CC BY-SA 4.0",
                              page="https://commons.wikimedia.org/wiki/File:Quincy_Washington_-_aerial.jpg"),
    "US at night, 2012": dict(credit="Image: NASA Scientific Visualization Studio, public domain", license="Public domain",
                              page="https://commons.wikimedia.org/wiki/File:Earth_at_Night_2012_(SVS30028_-_nightlights2012-united_states_4104).png"),
    "Liquid-cooled HPC data center, NREL ESIF": dict(
        credit="Photo: Dennis Schroeder, NREL / U.S. Department of Energy, public domain, via Wikimedia Commons",
        license="Public domain as a DOE posting (NREL is contractor-run; confirm before use)",
        page="https://commons.wikimedia.org/wiki/File:U.S._Department_of_Energy_-_Science_-_298_029_005_(31681202785).jpg"),
    "Priest Rapids Dam spillway, 1973": dict(
        credit="Photo: David Falconer, U.S. EPA DOCUMERICA / National Archives, public domain", license="Public domain",
        page="https://commons.wikimedia.org/wiki/File:THE_SPILLWAY_AT_PRIEST_RAPODS_DAM_-_NARA_-_548016.jpg"),
}


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------
def set_theme(prs):
    """Arial everywhere and the palette as theme colors, so new shapes match."""
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    for tag in ("majorFont", "minorFont"):
        for el in root.iter(f"{{{a}}}{tag}"):
            el.find(f"{{{a}}}latin").set("typeface", FONT)
    scheme = {"dk1": INK, "lt1": WHITE, "dk2": NAVY, "lt2": GROUND,
              "accent1": NAVY, "accent2": BLUE, "accent3": EMBER, "accent4": GREY,
              "accent5": RGBColor(0x14, 0x75, 0x8F), "accent6": GREEN}
    clr = root.find(f".//{{{a}}}clrScheme")
    clr.set("name", "Site engine")
    for name, rgb in scheme.items():
        el = clr.find(f"{{{a}}}{name}")
        for child in list(el):
            el.remove(child)
        etree.SubElement(el, f"{{{a}}}srgbClr", val=str(rgb))
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def style_run(run, size, color=INK, bold=False, italic=False):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color


def textbox(slide, x, y, w, h, text, size=20, color=INK, bold=False, italic=False,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None, wrap=True):
    """A real text box. `text` may be a string or a list of (text, overrides) runs per paragraph."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    paras = text if isinstance(text, list) else [text]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = para if isinstance(para, list) else [(para, {})]
        for t, o in runs:
            r = p.add_run()
            r.text = t
            style_run(r, o.get("size", size), o.get("color", color), o.get("bold", bold), o.get("italic", italic))
        if i > 0:
            p.space_before = Pt(6)
    return tb


def bullets(slide, x, y, w, h, items, size=22, color=INK, gap=14):
    """At most three bullets, each a real paragraph with a bullet character."""
    assert len(items) <= 3, "three bullets at most"
    tb = slide.shapes.add_textbox(x, y, w, h)
    tb.name = "Bullets"
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    for i, item in enumerate(items):
        assert len(item.split()) < 12 or item.startswith("[["), f"bullet over 11 words: {item}"
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.05
        p.space_after = Pt(gap)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(0.32)))
        pPr.set("indent", str(-Inches(0.32)))
        buClr = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(buClr, qn("a:srgbClr"), val=str(NAVY))
        etree.SubElement(pPr, qn("a:buSzPct"), val="100000")
        etree.SubElement(pPr, qn("a:buFont"), typeface=FONT)
        etree.SubElement(pPr, qn("a:buChar"), char="■")
        r = p.add_run()
        r.text = item
        style_run(r, size, color)
    return tb


def no_style(shape):
    """Drop the theme style reference, so no theme shadow or outline applies. Fill and line are explicit."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def hline(slide, x, y, w, color=RULE, weight=0.75, name="Rule"):
    """A hairline drawn as a thin filled rectangle, which renders the same everywhere."""
    return rect(slide, x, y, w, Emu(int(Pt(weight))), fill=color, name=name)


def rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=1.0, dash=False, name=None):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if name:
        s.name = name
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = 4  # MSO_LINE.DASH
    no_style(s)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.12
    return s


def shape_text(s, paras, size=16, color=INK, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.08):
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    paras = paras if isinstance(paras, list) else [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = para if isinstance(para, list) else [(para, {})]
        for t, o in runs:
            r = p.add_run()
            r.text = t
            style_run(r, o.get("size", size), o.get("color", color), o.get("bold", bold), o.get("italic", False))
    return s


def picture_fit(slide, path, x, y, w, h, align="center"):
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    px = x + (w - pw) // 2 if align == "center" else x
    py = y + (h - ph) // 2
    pic = slide.shapes.add_picture(str(path), px, py, pw, ph)
    return pic


def placeholder_box(slide, x, y, w, h, label, slide_id):
    s = rect(slide, x, y, w, h, fill=GROUND, line=INK3, dash=True, name="Placeholder")
    shape_text(s, label, size=16, color=INK2)
    PLACEHOLDERS.append((slide_id, label))
    return s


def bar_chart(slide, x, y, w, h, cats, vals, colors, title=None, fmt_code="#,##0", horizontal=True,
              max_val=None, cat_size=16, label_size=18, gap=45, labels=None):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("Value", vals)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(kind, x, y, w, h, cd)
    ch = gf.chart
    ch.has_legend = False
    ch.font.name = FONT
    ch.font.size = Pt(cat_size)
    ch.font.color.rgb = INK2
    if title:
        ch.has_title = True
        tf = ch.chart_title.text_frame
        tf.text = title
        for p in tf.paragraphs:
            p.alignment = PP_ALIGN.LEFT
            for r in p.runs:
                style_run(r, 18, INK, bold=True)
    else:
        ch.has_title = False
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.vary_by_categories = False
    ser = plot.series[0]
    for i, c in enumerate(colors):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = c
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = fmt_code
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(label_size)
    dl.font.bold = True
    dl.font.color.rgb = INK
    if labels:  # custom label text per point, still editable
        for i, t in enumerate(labels):
            if t is None:
                continue
            tf = ser.points[i].data_label.text_frame
            tf.text = t
            for r in tf.paragraphs[0].runs:
                style_run(r, label_size, INK, bold=True)
            ser.points[i].data_label.position = XL_LABEL_POSITION.OUTSIDE_END
    va = ch.value_axis
    va.has_major_gridlines = False
    va.visible = False
    va.minimum_scale = 0
    if max_val:
        va.maximum_scale = max_val
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(cat_size)
    ca.tick_labels.font.color.rgb = INK
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.format.line.color.rgb = RULE
    if horizontal:
        ca.reverse_order = True
    return ch


def rank_chart(slide, x, y, w, h, stages, series, title, max_rank):
    """Native line chart of ranks, #1 at the top. Each series: (name, ranks, color, labels, positions)."""
    pos = {"above": XL_LABEL_POSITION.ABOVE, "below": XL_LABEL_POSITION.BELOW,
           "left": XL_LABEL_POSITION.LEFT, "right": XL_LABEL_POSITION.RIGHT}
    cd = CategoryChartData()
    cd.categories = stages
    for name, ranks, *_ in series:
        cd.add_series(name, ranks)
    ch = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, x, y + Inches(0.45), w, h - Inches(0.45), cd).chart
    ch.has_legend = False
    ch.font.name = FONT
    ch.font.size = Pt(16)
    ch.has_title = False  # the title sits in a text box above, clear of the rank-1 labels
    textbox(slide, x, y, w, Inches(0.4), title, size=18, bold=True, align=PP_ALIGN.CENTER)
    va = ch.value_axis
    va.reverse_order = True
    va.crosses = XL_AXIS_CROSSES.MAXIMUM  # keeps the category axis at the bottom
    va.minimum_scale = -2  # headroom above rank 1 for its label
    va.maximum_scale = max_rank
    va.has_major_gridlines = False
    va.visible = False
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(15)
    ca.tick_labels.font.color.rgb = INK
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.format.line.color.rgb = RULE
    for ser, (name, ranks, color, labels, where) in zip(ch.plots[0].series, series):
        ser.smooth = False
        ser.format.line.color.rgb = color
        ser.format.line.width = Pt(4)
        ser.marker.style = XL_MARKER_STYLE.CIRCLE
        ser.marker.size = 13
        ser.marker.format.fill.solid()
        ser.marker.format.fill.fore_color.rgb = color
        ser.marker.format.line.color.rgb = color
        for i, (t, wpos) in enumerate(zip(labels, where)):
            dl = ser.points[i].data_label
            dl.text_frame.text = t
            for r in dl.text_frame.paragraphs[0].runs:
                style_run(r, 17, color, bold=True)
            dl.position = pos[wpos]
    return ch


def frame(prs, tag, headline, source, number, notes, checks=(), sources=()):
    """Blank slide with the section tag, headline, source line, rule, and slide number."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    rect(slide, MX, Inches(0.42), Inches(0.5), Inches(0.07), fill=NAVY, name="Accent")
    textbox(slide, MX, Inches(0.52), Inches(9), Inches(0.35), tag.upper(), size=13, color=NAVY, bold=True, name="Section tag")
    textbox(slide, MX, Inches(0.88), CW, Inches(1.2), headline, size=30, color=INK, bold=True, name="Headline",
            anchor=MSO_ANCHOR.TOP)
    hline(slide, MX, Inches(6.86), CW)
    textbox(slide, MX, Inches(6.92), CW - Inches(0.8), Inches(0.5), source, size=11, color=INK3, name="Source line")
    textbox(slide, W - MX - Inches(0.7), Inches(6.92), Inches(0.7), Inches(0.3), str(number), size=11, color=INK3,
            align=PP_ALIGN.RIGHT, name="Slide number")
    # Speaker notes: script, then every number with its file and field.
    script = notes.strip().split("\n\nIF ASKED")[0].split("\n\nALTERNATIVE CLOSE")[0]
    SCRIPT_WORDS[number] = len(script.replace("[Hand off to the live demo.]", "").split())
    lines = ["SCRIPT (about 30 to 35 seconds)", notes.strip(), ""]
    if checks:
        lines.append("NUMBERS ON THIS SLIDE ([[CHECK]] = could change tonight; recheck before presenting)")
        for what, value, f, field, check in checks:
            tag_ = " [[CHECK]]" if check else " (published source, stable)"
            lines.append(f"- {what}: {value}{tag_}. Source: {f}" + (f", field {field}" if field else ""))
            if check:
                CHECKS.append((number, what, value, f, field))
    if sources:
        lines.append("")
        lines.append("SOURCES")
        lines.extend(f"- {s}" for s in sources)
        SOURCES[number] = list(sources)
    slide.notes_slide.notes_text_frame.text = "\n".join(lines)
    return slide


def quote_block(slide, x, y, w, h, q, slide_id, topic, size=24):
    """A verified quote with attribution, or a labeled placeholder."""
    if not q:
        return placeholder_box(slide, x, y, w, h, f"[[QUOTE NEEDED: {topic}]]", slide_id)
    s = rect(slide, x, y, w, h, fill=GROUND, name="Quote")
    bar = rect(slide, x, y, Inches(0.08), h, fill=NAVY, name="Quote bar")
    tb = textbox(slide, x + Inches(0.3), y + Inches(0.12), w - Inches(0.5), h - Inches(0.2),
                 [[("“" + q["text"] + "”", {"size": size, "italic": False, "bold": False})],
                  [(q.get("cite", q["who"]), {"size": 14, "color": INK2, "bold": True})]],
                 size=size, color=INK, name="Quote text", anchor=MSO_ANCHOR.MIDDLE)
    return tb


# ---------------------------------------------------------------------------
# Data
# ---------------------------------------------------------------------------
F = FACTS
B = F["balanced"]
TOP10 = F["top10"]
GRANT = F["featured"]
GR = F["grant_ranges"]
IMP = F["impact"]
W8 = F["weights"]
PS = F["pick_story"]
GL = F["global"]

SMAA = []
with open(ROOT / "scratch/weighting/out/smaa_acceptability.csv") as fh:
    for row in csv.DictReader(fh):
        SMAA.append(row)
SMAA_TOP10 = sorted(SMAA, key=lambda r: -float(r["top10_floor_on"]))[:6]
SMAA_RANK1 = sorted(SMAA, key=lambda r: -float(r["rank1_floor_on"]))[:3]


def smaa(fips, col):
    return float(next(r for r in SMAA if r["fips"] == fips)[col])


# Typed-in figures from research write-ups (file named next to each).
PUD = {  # research/risk.md and research/implementation.md, from the Grant PUD FAQ (2026-08-28)
    "hydro_share_amw": 633, "load_2025_amw": 757, "campus_amw": 279, "queued_mw": 800}
TIE = {  # docs/weighting.md, Recommendation table, 25-year cost at $190/t with sales tax
    "Clark, WA": 4.405, "Grant, WA": 4.452, "Franklin, NY": 4.454}
PRICE_RERUN = {80: 79, 132: 1146}  # deck build log step 1: Grant's rank with WA at BPA's new-load rate
EVAP = {"passed": 826, "floor_ok": 542, "first": "Wayne, TN"}  # docs/demo_script.md, rechecked in the build log
PCTL = {"Grant, WA": (70.4, 93.6), "Franklin, NY": (80.4, 98.7)}  # build log step 1: pillar score, national pctl
FRANKLIN_CO2 = 262168  # research/impact.md, dry cooling table


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    set_theme(prs)
    slide_hook(prs)
    slide_answer(prs)
    slide_framework(prs)
    slide_impact(prs)
    slide_pick(prs)
    slide_robust(prs)
    slide_risk(prs)
    slide_plan(prs)
    slide_engine(prs)
    app_tonnes(prs)
    app_tie(prs)
    app_weights(prs)
    app_sources(prs)
    app_limits(prs)
    app_extends(prs)
    out = DECK / "pitch_template.pptx"
    prs.save(out)
    write_checks()
    write_sources()
    print(f"wrote {out.relative_to(ROOT)}: {len(prs.slides)} slides, {len(CHECKS)} numbers to check, "
          f"{len(PLACEHOLDERS)} placeholders")
    for s, t in PLACEHOLDERS:
        print(f"  placeholder on {s}: {t}")
    print("  script words (about 150 a minute, so 75 to 90 for 30 to 35 s):",
          ", ".join(f"{k}: {v}" for k, v in SCRIPT_WORDS.items()))


# ---------------------------------------------------------------------------
# Main slides
# ---------------------------------------------------------------------------
def slide_hook(prs):
    q = QUOTES["ny_eo62"]
    s = frame(
        prs, "Why siting matters",
        "Where AI campuses go now locks in decades of carbon, water, and cost.",
        "Sources: LBNL, 2024 United States Data Center Energy Usage Report (Dec. 2024), Executive Summary; "
        "New York Executive Order No. 62 (July 14, 2026), WHEREAS clause 9; research/impact.md. Links in docs/deck/sources.md.",
        1,
        notes="""
AI is driving a data center boom. Berkeley Lab found data centers used 4.4 percent of
US electricity in 2023 and projects 6.7 to 12 percent by 2028. Communities are pushing
back on where they go: New York paused state permits for large data centers this summer,
citing energy and water. A campus built now runs 20 to 30 years, so the county you pick
sets its carbon, its water, and its power bill for decades. Grid carbon alone varies
about four times across the counties we scored.
""",
        checks=[("US data center share of electricity, 2023", "4.4% (176 TWh)", "LBNL 2024 report, Executive Summary p. 5 (sources.md)", None, False),
                ("Projected share, 2028", "6.7% to 12.0%", "LBNL 2024 report, Executive Summary p. 6 (sources.md)", None, False),
                ("Grid carbon varies about 4x across candidate counties", "4x (242 to 911 lb/MWh)",
                 "research/impact.md", "What the numbers say; dry cooling table", True)],
        sources=[f"LBNL 2024 United States Data Center Energy Usage Report (Shehabi et al.), {QUOTES['lbnl_2023']['url']}",
                 f"New York Executive Order No. 62, {q['url']}"],
    )
    quote_block(s, MX, TOP, Inches(5.7), Inches(2.45), q, 1, "policy pushback on data center siting", size=21)
    bullets(s, MX, TOP + Inches(2.7), Inches(5.7), Inches(1.8),
            ["A campus built now runs 20 to 30 years.",
             "Grid carbon varies about 4x across candidate counties."], size=21, gap=10)
    bar_chart(s, MX + Inches(6.1), TOP, Inches(6.0), Inches(4.5),
              ["2023", "2028, low", "2028, high"], [4.4, 6.7, 12.0], [NAVY, BLUE, BLUE],
              title="Data centers' share of US electricity (%)", fmt_code='0.0"%"', horizontal=False,
              max_val=14, cat_size=18, label_size=22, gap=60)


def slide_answer(prs):
    g = GRANT
    s = frame(
        prs, "The answer",
        "Grant County, Washington, if the campus funds its own clean power.",
        "Sources: team engine, balanced preset (results/balanced.csv; docs/figures/facts.json); "
        "map docs/figures/map_composite.png; 25-year cost docs/weighting.md; Grant PUD data center FAQ, 2026-08-28.",
        2,
        notes=f"""
Our answer is Grant County, Washington, home of the Quincy data center cluster. Of
{fmt(B['counties'])} counties, {fmt(B['passed'])} pass every hard gate, and Grant leads them at
today's average power prices. It makes the top 10 under half of all possible weightings,
more than any other county. In 25-year dollars it ties Clark, Washington and Franklin, New
York; we feature Grant because its power timeline and tax status are sourced. The condition:
sited next to hydro, powered by new clean supply the project funds.
""",
        checks=[("Counties scored", fmt(B["counties"]), "docs/figures/facts.json", "balanced.counties", True),
                ("Counties passing the gates", fmt(B["passed"]), "docs/figures/facts.json", "balanced.passed", True),
                ("Grant rank and composite", f"#{g['rank']}, {g['composite']}", "docs/figures/facts.json", "featured.rank, featured.composite", True),
                ("Lead over #2", f"{F['gap_first_to_second']} points", "docs/figures/facts.json", "gap_first_to_second", True),
                ("Grant top-10 share over 5,000 random weightings", pct(smaa('53025', 'top10_floor_on')),
                 "scratch/weighting/out/smaa_acceptability.csv", "top10_floor_on", True),
                ("Three-county 25-year cost spread", "within 1.1%", "docs/weighting.md", "Recommendation", True)],
        sources=["Grant PUD data center FAQ, 2026-08-28, https://www.grantpud.org/blog/data-center-faqs"],
    )
    picture_fit(s, FIG / "map_composite.png", MX, TOP - Inches(0.05), Inches(6.4), Inches(4.8))
    x = MX + Inches(6.75)
    w = CW - Inches(6.75)
    box = rect(s, x, TOP, w, Inches(1.55), fill=NAVY, name="Condition")
    shape_text(box, [[("THE CONDITION", {"size": 13, "bold": True, "color": PALE})],
                     [("Sited next to hydro, powered by new clean supply the project funds.",
                       {"size": 21, "bold": True, "color": WHITE})]],
               align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=0.2)
    bullets(s, x, TOP + Inches(1.85), w, Inches(2.9),
            [f"Leads {fmt(B['passed'])} gate-passing counties at today's average power prices.",
             f"Top 10 in {pct(smaa('53025', 'top10_floor_on'), 0)} of all weightings, more than any county.",
             "Ties Clark WA and Franklin NY within 1.1% over 25 years."], size=20, gap=12)


def slide_framework(prs):
    s = frame(
        prs, "How it decides",
        f"Gates cut {fmt(B['counties'])} counties to {fmt(B['passed'])}. Eight pillars rank the rest.",
        "Sources: engine/conditions/balanced.yaml; engine/pillars.yaml; docs/figures/facts.json; docs/weighting.md. "
        "Data from 20 public sources, listed in appendix A4.",
        3,
        notes=f"""
Here's how the engine decides. Hard gates come first: flood, wildfire, fiber, queue age,
moratoria, and enough generation nearby. {fmt(B['passed'])} counties pass. Eight pillars then
score each county against the nation, with stated weights. A county weak on any pillar ranks
below every county that isn't; {fmt(B['floor_ok'])} clear that floor. The ranking isn't machine
learning. For the leaders, a 25-year dollar model prices energy, carbon, water, hazard,
delay, and tax, and tells us what must be true for each one to win.
""",
        checks=[("Counties", fmt(B["counties"]), "docs/figures/facts.json", "balanced.counties", True),
                ("Pass the gates", fmt(B["passed"]), "docs/figures/facts.json", "balanced.passed", True),
                ("Clear the pillar floor", fmt(B["floor_ok"]), "docs/figures/facts.json", "balanced.floor_ok", True),
                ("Pillar weights", ", ".join(f"{k} {v:.3f}" for k, v in W8["balanced"].items()),
                 "docs/figures/facts.json", "weights.balanced (from engine/conditions/balanced.yaml)", True)],
    )
    steps = [
        (fmt(B["counties"]), "contiguous US counties"),
        (fmt(B["passed"]), "pass every hard gate"),
        (fmt(B["floor_ok"]), "clear the pillar floor"),
        ("$ and t", "25-year cost and carbon price the leaders"),
        ("Research", "power, water, and tax decide"),
    ]
    n = len(steps)
    gap = Inches(0.42)
    bw = int((CW - gap * (n - 1)) / n)
    bh = Inches(1.75)
    y = TOP + Inches(0.1)
    for i, (big, small) in enumerate(steps):
        x = MX + i * (bw + gap)
        fill = NAVY if i < 3 else BLUE
        b = rect(s, x, y, bw, bh, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=f"Step {i + 1}")
        big_size = 32 if big[0].isdigit() else 24
        shape_text(b, [[(big, {"size": big_size, "bold": True, "color": WHITE})],
                       [(small, {"size": 16, "color": WHITE})]])
        if i < n - 1:
            a = rect(s, x + bw + Inches(0.07), y + bh / 2 - Inches(0.16), gap - Inches(0.14), Inches(0.32),
                     fill=INK3, shape=MSO_SHAPE.RIGHT_ARROW, name="Arrow")
    textbox(s, MX, y + bh + Inches(0.35), CW, Inches(0.4),
            "Eight pillars, each a national percentile, weighted (balanced preset):", size=16, color=INK2, bold=True)
    cy = y + bh + Inches(0.85)
    cg = Inches(0.1)
    cwid = int((CW - cg * 7) / 8)
    for i, (key, label, hue) in enumerate(PILLARS):
        c = rect(s, MX + i * (cwid + cg), cy, cwid, Inches(0.95), fill=hue, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                 name=f"Pillar {label}")
        shape_text(c, [[(label, {"size": 14, "bold": True, "color": WHITE})],
                       [(f"{W8['balanced'][key] * 100:.0f}%", {"size": 20, "bold": True, "color": WHITE})]], margin=0.04)
    textbox(s, MX, cy + Inches(1.1), CW, Inches(0.4),
            "Not machine learning: stated weights, tested for robustness on slide 6.", size=15, color=INK3)


def slide_impact(prs):
    co2_bpa = GR["co2_tonnes"]["bpa"]
    co2_nw = GR["co2_tonnes"]["nwpp_table"]
    lou = IMP["Loudoun, VA"]
    gi = IMP["Grant, WA"]
    cut = 1 - gi["water_million_gal_dry"] / gi["water_million_gal_evap"]
    s = frame(
        prs, "Sustainability impact",
        f"Dry cooling cuts water {cut * 100:.0f}%. State law forces power clean by 2045.",
        "Sources: etl/impact.py and research/impact.md (300 MW IT, load factor 0.8; PUE and WUE from Lei and Masanet); "
        "eGRID2023; RCW 19.405; Washington Data Center Workgroup Preliminary Report (Dec. 2025), Finding 19, p. 13.",
        4,
        notes=f"""
Here's the sustainability case. Water first: an evaporative campus in Grant would use about
{gi['water_million_gal_evap']:.0f} million gallons a year. Our plan uses dry cooling, about
{gi['water_million_gal_dry']:.0f} million, an {cut * 100:.0f} percent cut, for about 2 percent more
energy. Washington's own data center workgroup flags water as a real concern. Carbon depends
on which power serves us. Before our funded supply arrives, it's {co2_bpa / 1000:.0f} to
{co2_nw / 1000:.0f} thousand tons a year, against Loudoun's {lou['co2_tonnes_dry'] / 1000:.0f}
thousand. State law then requires the utility's power to be fully clean by 2045.

IF ASKED:
- Why not zero carbon? The utility's hydro rate of 0 lb/MWh describes its existing customers.
  None of that hydro is spare for a new load, so we don't claim it.
- Why compare with Loudoun? It's where the industry builds today. It fails our queue-age gate,
  so it's a reference, not a candidate.
- State law detail: utility sales must be greenhouse-gas neutral by 2030 (up to 20% through
  alternative compliance) and 100% renewable or non-emitting by 2045 (RCW 19.405.040, .050).
""",
        checks=[("Grant water, evaporative", f"{gi['water_million_gal_evap']} M gal/yr", "docs/figures/facts.json", "impact['Grant, WA'].water_million_gal_evap", True),
                ("Grant water, dry", f"{gi['water_million_gal_dry']} M gal/yr", "docs/figures/facts.json", "impact['Grant, WA'].water_million_gal_dry", True),
                ("Loudoun water, evaporative", f"{lou['water_million_gal_evap']} M gal/yr", "docs/figures/facts.json", "impact['Loudoun, VA'].water_million_gal_evap", True),
                ("Water cut, dry vs evaporative", f"{cut * 100:.0f}%", "computed from the two Grant values above", None, True),
                ("Grant CO2 at BPA's mix", f"{fmt(co2_bpa)} t/yr", "docs/figures/facts.json", "grant_ranges.co2_tonnes.bpa", True),
                ("Grant CO2 at the Northwest average", f"{fmt(co2_nw)} t/yr", "docs/figures/facts.json", "grant_ranges.co2_tonnes.nwpp_table", True),
                ("Loudoun CO2, dry", f"{fmt(lou['co2_tonnes_dry'])} t/yr", "docs/figures/facts.json", "impact['Loudoun, VA'].co2_tonnes_dry", True),
                ("Dry cooling energy penalty", "about 2%", "research/risk.md", "Water stress row", True),
                ("CETA dates", "2030 neutral, 2045 100% clean", "research/implementation.md", "Carbon over 30 years", False)],
        sources=["Washington Clean Energy Transformation Act, RCW 19.405.040 and 19.405.050, https://app.leg.wa.gov/RCW/default.aspx?cite=19.405",
                 f"Washington Data Center Workgroup, Preliminary Report, Finding 19, p. 13, {QUOTES['wa_water']['url']}"],
    )
    half = int((CW - Inches(0.5)) / 2)
    bar_chart(s, MX, TOP - Inches(0.05), half, Inches(3.15),
              ["Loudoun VA, evaporative", "Grant, evaporative", "Grant, dry (+2% energy)"],
              [lou["water_million_gal_evap"], gi["water_million_gal_evap"], gi["water_million_gal_dry"]],
              [GREY, BLUE, NAVY], title="On-site water, million gallons a year", max_val=400, cat_size=15)
    bar_chart(s, MX + half + Inches(0.5), TOP - Inches(0.05), half, Inches(3.15),
              ["Loudoun VA", "Grant, Northwest average", "Grant, BPA's mix"],
              [lou["co2_tonnes_dry"] / 1000, co2_nw / 1000, co2_bpa / 1000],
              [GREY, BLUE, NAVY], title="CO2, thousand metric tons a year", max_val=850, cat_size=15)
    quote_block(s, MX, TOP + Inches(3.25), half, Inches(1.25), QUOTES["wa_water"], 4, "water use concerns", size=17)
    bullets(s, MX + half + Inches(0.5), TOP + Inches(3.35), half, Inches(1.1),
            ["Shown before funded supply. State law: 100% clean by 2045."], size=19)


def slide_pick(prs):
    r = PS["ranks"]
    s = frame(
        prs, "How the pick changed",
        "Corrections moved our pick. The last one cut against Grant.",
        "Sources: docs/figures/facts.json (pick_story; same data as docs/figures/pick_story.png); research/implementation.md (BPA new-load rate, "
        "BP-26 rate schedules); research/permitting_model.md; data/processed/permitting_validation.json.",
        5,
        notes=f"""
We didn't start at Grant. With seven pillars, Berkshire County, Massachusetts ranked first
and Grant seventh. Massachusetts closed the tax exemption Berkshire relied on, and it fell
to fifth. Then we made power cost its own pillar. Grant
rose to first, and Berkshire fell to {fmt(r['Berkshire, MA'][2])}th. The last correction cut
against us. A new 300 megawatt load doesn't get today's average price, so Grant's bill is 196
to 323 million dollars a year, not 162. And we dropped a permitting model that failed out of sample.

IF ASKED: why doesn't Grant get cheap hydro? BPA's rule for a new large single load: "{QUOTES['bpa_nlsl']['text']}"
({QUOTES['bpa_nlsl']['who']}). Whether the rule binds a Grant PUD load depends on the utility's BPA contract; not checked.
""",
        checks=[("Berkshire ranks by stage", " / ".join(map(str, r["Berkshire, MA"])), "docs/figures/facts.json", "pick_story.ranks['Berkshire, MA']", True),
                ("Grant ranks by stage", " / ".join(map(str, r["Grant, WA"])), "docs/figures/facts.json", "pick_story.ranks['Grant, WA']", True),
                ("Grant energy cost at state average", f"${GR['energy_cost_musd']['state_average']}M/yr", "docs/figures/facts.json", "grant_ranges.energy_cost_musd.state_average", True),
                ("Grant energy cost at BPA new-load rate", f"${GR['energy_cost_musd']['new_load_low']}M to ${GR['energy_cost_musd']['new_load_high']}M/yr",
                 "docs/figures/facts.json", "grant_ranges.energy_cost_musd.new_load_low/high", True),
                ("Permitting model AUC without facility counts", "0.48 to 0.585 (bar 0.60)", "research/permitting_model.md; data/processed/permitting_validation.json", None, True)],
    )
    g, bk = r["Grant, WA"], r["Berkshire, MA"]
    off = 25  # Berkshire's last rank (1,047) is drawn at the bottom edge and labeled
    rank_chart(s, MX, TOP - Inches(0.1), Inches(6.9), Inches(4.75),
               ["Seven pillars", "MA exemption closed", "Cost as a pillar"],
               [("Berkshire, MA", [bk[0], bk[1], min(bk[2], off)], EMBER,
                 [f"Berkshire #{bk[0]}", f"#{bk[1]}", f"#{fmt(bk[2])}"], ["above", "above", "left"]),
                ("Grant, WA", g, NAVY, [f"Grant #{g[0]}", f"#{g[1]}", f"Grant #{g[2]}"], ["below", "below", "above"])],
               title="Balanced-preset rank at each step", max_rank=off + 1)
    x = MX + Inches(7.2)
    w = CW - Inches(7.2)
    bullets(s, x, TOP + Inches(0.1), w, Inches(3.2),
            ["Massachusetts closed a tax exemption: Berkshire fell to 5th.",
             f"Power cost as its own pillar: Grant {r['Grant, WA'][1]}th to 1st.",
             f"New-load power prices raise Grant's bill to ${GR['energy_cost_musd']['new_load_low']}–{GR['energy_cost_musd']['new_load_high']}M."],
            size=20, gap=14)
    box = rect(s, x, TOP + Inches(3.35), w, Inches(1.3), fill=GROUND, name="Dropped")
    shape_text(box, [[("TESTED AND DROPPED", {"size": 13, "bold": True, "color": EMBER})],
                     [("Permitting ML model: out-of-sample AUC 0.48–0.59, below our 0.60 bar.", {"size": 17})]],
               align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=0.18)


def slide_robust(prs):
    names = [r["county"] for r in SMAA_TOP10]
    vals = [float(r["top10_floor_on"]) * 100 for r in SMAA_TOP10]
    colors = [NAVY if r["fips"] == "53025" else BLUE for r in SMAA_TOP10]
    r1 = SMAA_RANK1
    s = frame(
        prs, "Robustness",
        "Grant makes the top 10 under more weightings than any county.",
        "Sources: SMAA over the engine's eight pillars, 5,000 Dirichlet weight draws, floor on "
        "(scratch/weighting/out/smaa_acceptability.csv; docs/weighting.md); results/balanced.csv (robustness).",
        6,
        notes=f"""
The obvious question: did we pick weights that make Grant win? So we stopped choosing. We
drew 5,000 random weightings of the eight pillars and counted how often each county makes
the top 10. Grant does in {pct(smaa('53025', 'top10_floor_on'), 1)} of them, more than any
other county. No county wins most weightings; the top three each come first in 11 to 13
percent. Near our own weights, Grant stays in the top 10 in {pct(GRANT['robustness'], 1)}.
We didn't tune the weights to get our answer. We tested whether our answer depends on them.
""",
        checks=[(f"SMAA top-10 share, {n}", f"{v:.1f}%", "scratch/weighting/out/smaa_acceptability.csv", "top10_floor_on", True)
                for n, v in zip(names, vals)]
        + [(f"SMAA rank-1 share, {r['county']}", pct(float(r['rank1_floor_on'])), "scratch/weighting/out/smaa_acceptability.csv", "rank1_floor_on", True) for r in r1]
        + [("Robustness near balanced weights", pct(GRANT["robustness"]), "docs/figures/facts.json", "featured.robustness (2,000 draws, sd about 0.03)", True)],
    )
    bar_chart(s, MX, TOP, Inches(6.6), Inches(4.0), names, vals, colors,
              title="Top 10 in this share of random weightings", fmt_code='0.0"%"', max_val=65, cat_size=16)
    x = MX + Inches(6.95)
    w = CW - Inches(6.95)
    lo = min(float(r["rank1_floor_on"]) for r in r1) * 100
    hi = max(float(r["rank1_floor_on"]) for r in r1) * 100
    bullets(s, x, TOP + Inches(0.1), w, Inches(3.0),
            ["5,000 random weightings: we chose none of them.",
             f"No county wins most: the top three each win {lo:.0f}–{hi:.0f}%.",
             f"Near our weights, Grant stays top 10 in {pct(GRANT['robustness'])}."],
            size=20, gap=14)
    band = rect(s, MX, TOP + Inches(4.15), CW, Inches(0.6), fill=GROUND, name="Line")
    shape_text(band, "We didn't tune the weights to get our answer. We tested whether our answer depends on them.",
               size=19, bold=True, color=NAVY)


def slide_risk(prs):
    p = PUD
    s = frame(
        prs, "Risk and win condition",
        "Grant's case rests on one thing: full power by about 2029.",
        "Sources: Grant PUD data center FAQ (2026-08-28); research/risk.md; docs/weighting.md; protected land and tribal rows "
        "from research/risk.md on branch fix/sensitive-land (d2ed911, not yet on main; PAD-US 4.1).",
        7,
        notes=f"""
One risk is high: power. Grant PUD's share of its dams averages about
{p['hydro_share_amw']} average megawatts, already below its {p['load_2025_amw']} load. We'd add
{p['campus_amw']}, and about {p['queued_mw']} megawatts are queued ahead of us. So we
fund new supply and phase in behind the 2027 and 2029 transmission lines. Timing is the win
condition: in 25-year dollars Grant ties Clark and Franklin, and each month of delay costs
about 25 million dollars. Heat is a medium risk. Protected land and tribal sites are low,
checked by hand, with early consultation.

IF ASKED (not in the 35 seconds):
- Price: with Washington at BPA's new-load rate and other states at their averages, Grant
  falls to {PRICE_RERUN[80]}th at $80/MWh and {fmt(PRICE_RERUN[132])}th at $132/MWh. No other state's
  new-load rate is sourced, so that comparison is lopsided, but Grant's #1 rests on today's
  average price.
- Hydro and fisheries: Washington's Data Center Workgroup (Preliminary Report, Finding 19c,
  p. 13) says new load on hydropower increases demand for a scarce resource tied to Tribal and
  state fisheries efforts. Answer: we don't claim existing hydro; we fund new supply.
- Regional context (verified quote): "{QUOTES['wa_load_growth']['text']}" {QUOTES['wa_load_growth']['who']}.
""",
        checks=[("Grant PUD share of its dams", f"about {p['hydro_share_amw']} aMW", "research/risk.md (Grant PUD FAQ)", "Power availability row", False),
                ("Grant PUD 2025 load", f"{p['load_2025_amw']} aMW", "research/risk.md (Grant PUD FAQ)", "Power availability row", False),
                ("Campus average load", f"{p['campus_amw']} aMW", "research/risk.md; etl/impact.py", "Power availability row", True),
                ("Large-load requests queued", f"about {p['queued_mw']} MW", "research/risk.md (Grant PUD FAQ)", "Power availability row", False),
                ("Delay cost", "$25M per month", "docs/weighting.md", "Assumptions: time to power", True),
                ("Grant break-even vs Clark / Franklin", "10.2 / 12.1 months past a 2-year baseline", "docs/weighting.md", "Grant's row", True),
                ("Days above 95°F", f"{GRANT['horizon_2050_raw']['days_above_95f_hist']['today']:.1f} today, "
                 f"{GRANT['horizon_2050_raw']['days_above_95f_hist']['days_above_95f_2050_rcp85']:.1f} by 2050",
                 "docs/figures/facts.json", "featured.horizon_2050_raw.days_above_95f_hist", True),
                ("Grant rank with WA at BPA's $80 and $132/MWh", f"{PRICE_RERUN[80]} and {fmt(PRICE_RERUN[132])}",
                 "docs/deck_build_log.md", "step 1, price sensitivity (scratchpad rerun, not committed)", True),
                ("Protected land and tribal rows", "residual Low", "research/risk.md on origin/fix/sensitive-land (d2ed911)", "Protected and sensitive land; Tribal consultation", True)],
        sources=["Grant PUD data center FAQ, 2026-08-28, https://www.grantpud.org/blog/data-center-faqs",
                 "PAD-US 4.1 and tribal land checks: research/risk.md and research/sensitive_land.md on branch fix/sensitive-land, commit d2ed911",
                 f"Washington Data Center Workgroup, Preliminary Report, Findings 6 and 19c, {QUOTES['wa_load_growth']['url']}"],
    )
    bar_chart(s, MX, TOP - Inches(0.05), Inches(5.6), Inches(2.45),
              ["Utility's share of its dams", "Utility load, 2025", "Load plus our campus"],
              [p["hydro_share_amw"], p["load_2025_amw"], p["load_2025_amw"] + p["campus_amw"]],
              [BLUE, GREY, EMBER], title="Grant PUD, average megawatts", max_val=1300, cat_size=15, label_size=17, gap=35)
    quote_block(s, MX, TOP + Inches(2.5), Inches(5.6), Inches(1.0), QUOTES["pud_speed"], 7, "Grant PUD on new large loads", size=16)
    x = MX + Inches(5.95)
    w = CW - Inches(5.95)
    rows = [("HIGH", EMBER, f"Power: no spare hydro; ~{p['queued_mw']} MW of requests queued ahead."),
            ("MEDIUM", AMBER, "Heat: days above 95°F rise from 14 to 36 by 2050."),
            ("LOW", GREEN, "Protected land, tribal sites: checked by hand; consult early.")]
    for i, (lvl, c, text) in enumerate(rows):
        y = TOP + i * Inches(1.17)
        chip = rect(s, x, y + Inches(0.2), Inches(1.2), Inches(0.5), fill=c, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=f"Risk {lvl}")
        shape_text(chip, lvl, size=14, bold=True, color=WHITE)
        textbox(s, x + Inches(1.4), y, w - Inches(1.4), Inches(0.95), text, size=19, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    box = rect(s, MX, TOP + Inches(3.65), CW, Inches(0.9), fill=NAVY, name="Win condition")
    shape_text(box, [[("WIN CONDITION  ", {"size": 14, "bold": True, "color": PALE}),
                      ("In 25-year dollars Grant ties Clark WA and Franklin NY. Each month of power delay costs about $25M.",
                       {"size": 19, "bold": True, "color": WHITE})]],
               align=PP_ALIGN.LEFT, margin=0.25)


def slide_plan(prs):
    s = frame(
        prs, "The 30-year plan",
        "Phase in behind new transmission, fund new clean supply, use no evaporative water.",
        "Sources: research/implementation.md (Columbia Basin Herald 2025-03-31, Rye Development, RCW 19.405, ASHRAE "
        "liquid cooling classes, Ramboll Odense case); Grant PUD Data Center FAQs (Aug. 28, 2026), Q5.",
        8,
        notes="""
The plan has three parts. Power: start small behind the utility's queue, energize after
Quincy's 2027 upgrade, and grow toward 300 megawatts after the 2029 line, only as supply we
pay for comes online. About a gigawatt of new solar matches our annual use, not every hour,
so firm clean power follows in the 2030s. Cooling: closed-loop liquid cooling with dry
coolers, sized for 2050's hotter days. Heat: we site next to a food processor and pipe 45 to
65 degree return water into its process heat.
""",
        checks=[("Transmission milestones", "2027 Quincy upgrade; 2029 Wanapum to Quincy line", "research/implementation.md", "The queue and the caps", False),
                ("New solar to match annual use", "about 1 GW (1,030 MW at 27% capacity factor)", "research/implementation.md", "Where new clean supply comes from", True),
                ("Goldendale pumped storage", "1.2 GW, 2031 to 2032", "research/implementation.md", "Where new clean supply comes from", False),
                ("Dry cooling water", f"{IMP['Grant, WA']['water_million_gal_dry']} M gal/yr", "docs/figures/facts.json", "impact['Grant, WA'].water_million_gal_dry", True),
                ("Liquid-cooling return water", "45 to 65°C", "research/implementation.md", "Cooling and water; Heat reuse", False)],
        sources=[f"Grant PUD Data Center FAQs, Q5, {QUOTES['pud_pays']['url']}"],
    )
    # Timeline across the top half: native shapes, all editable.
    ms = [("2027", "Quincy transmission upgrade. Phase 1 energizes."),
          ("2029", "Wanapum–Quincy line. Grow toward 300 MW."),
          ("2030", "State law: utility sales GHG-neutral."),
          ("2030s", "Firm clean supply: pumped storage."),
          ("2045", "State law: 100% clean retail power.")]
    y_line = TOP + Inches(0.65)
    hline(s, MX + Inches(0.3), y_line - Emu(int(Pt(1.5))), CW - Inches(0.6), color=NAVY, weight=3, name="Timeline")
    slot = int(CW / len(ms))
    for i, (yr, text) in enumerate(ms):
        cx = MX + i * slot + slot // 2
        rect(s, cx - Inches(0.14), y_line - Inches(0.14), Inches(0.28), Inches(0.28), fill=NAVY if i < 2 else BLUE,
             shape=MSO_SHAPE.OVAL, name=f"Milestone {yr}")
        textbox(s, cx - slot // 2, TOP - Inches(0.1), slot, Inches(0.5), yr, size=24, bold=True, color=NAVY,
                align=PP_ALIGN.CENTER)
        textbox(s, cx - slot // 2 + Inches(0.08), y_line + Inches(0.22), slot - Inches(0.16), Inches(0.85), text,
                size=16, color=INK, align=PP_ALIGN.CENTER)
    # Three cards: power, cooling, heat. One bullet each.
    cards = [("POWER", "About 1 GW of new solar matches annual use, not hourly."),
             ("COOLING", f"Closed-loop liquid, dry coolers: {IMP['Grant, WA']['water_million_gal_dry']:.0f}M gallons a year."),
             ("HEAT REUSE", "45–65°C return water feeds a neighboring food processor.")]
    cg = Inches(0.3)
    cwid = int((CW - cg * 2) / 3)
    cy = TOP + Inches(1.85)
    for i, (head, text) in enumerate(cards):
        assert len(text.split()) < 12, text
        c = rect(s, MX + i * (cwid + cg), cy, cwid, Inches(1.5), fill=GROUND, name=f"Card {head}")
        rect(s, MX + i * (cwid + cg), cy, cwid, Inches(0.07), fill=NAVY, name="Card rule")
        shape_text(c, [[(head, {"size": 14, "bold": True, "color": NAVY})], [(text, {"size": 19})]],
                   align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=0.22)
    quote_block(s, MX, TOP + Inches(3.5), CW, Inches(1.0), QUOTES["pud_pays"], 8, "Grant PUD on who pays for new supply", size=17)


def slide_engine(prs):
    q = None  # close on the team line; a verified alternative is in the notes
    s = frame(
        prs, "The engine is the product",
        "Change one condition and the engine gives a new, explained answer.",
        "Sources: engine/conditions/balanced.yaml with cooling: evaporative (rerun in docs/deck_build_log.md step 1); "
        "docs/demo_script.md; WRI Aqueduct 4.0 water stress.",
        9,
        notes=f"""
The deliverable isn't one county. It's the engine. Conditions go in as one plain file,
and a ranked, explained shortlist comes out. Watch one switch: change cooling from dry to
evaporative. The engine adds a water-stress gate, and the counties that pass drop from
{fmt(B['passed'])} to {fmt(EVAP['passed'])}. Grant scores 3.6 of 5 for water stress, so it drops
out, and the engine says exactly why. Wayne, Tennessee leads instead. Same data, new
answer, with reasons. Let's show you. [Hand off to the live demo.]

ALTERNATIVE CLOSE (verified quote): "{QUOTES['wa_siting']['text']}" {QUOTES['wa_siting']['who']}.
""",
        checks=[("Counties passing, dry", fmt(B["passed"]), "docs/figures/facts.json", "balanced.passed", True),
                ("Counties passing, evaporative", fmt(EVAP["passed"]), "docs/deck_build_log.md step 1; docs/demo_script.md", "0:2:00 row", True),
                ("Grant water stress", f"{GRANT['horizon_2050_raw']['water_stress_bws']['today']:.1f} of 5", "docs/figures/facts.json", "featured.horizon_2050_raw.water_stress_bws.today", True),
                ("New leader under evaporative cooling", EVAP["first"], "docs/deck_build_log.md step 1", None, True)],
    )
    bar_chart(s, MX, TOP, Inches(5.6), Inches(3.3), ["Dry cooling", "Evaporative cooling"],
              [B["passed"], EVAP["passed"]], [NAVY, BLUE], title="Counties that pass the gates", max_val=2000, cat_size=17)
    x = MX + Inches(5.95)
    w = CW - Inches(5.95)
    bullets(s, x, TOP + Inches(0.1), w, Inches(2.9),
            ["Evaporative cooling switches on a water-stress gate.",
             f"Grant (water stress {GRANT['horizon_2050_raw']['water_stress_bws']['today']:.1f} of 5) drops out.",
             "One conditions file in; ranked, explained shortlist out."],
            size=20, gap=14)
    band = rect(s, MX, TOP + Inches(3.55), CW, Inches(1.15), fill=NAVY, name="Hand-off")
    if q:
        shape_text(band, [[("“" + q["text"] + "”  ", {"size": 20, "bold": True, "color": WHITE}),
                           (q.get("cite", q["who"]), {"size": 15, "color": PALE})]], align=PP_ALIGN.LEFT, margin=0.3)
    else:
        shape_text(band, [[("LIVE DEMO  →  ", {"size": 16, "bold": True, "color": PALE}),
                           ("Conditions in. A ranked, explained shortlist out.", {"size": 24, "bold": True, "color": WHITE})]],
                   align=PP_ALIGN.LEFT, margin=0.3)


# ---------------------------------------------------------------------------
# Appendix
# ---------------------------------------------------------------------------
def app_tonnes(prs):
    g_score, g_pct = PCTL["Grant, WA"]
    f_score, f_pct = PCTL["Franklin, NY"]
    co2_g = GR["co2_tonnes"]["nwpp_table"]
    ratio = co2_g / FRANKLIN_CO2
    s = frame(
        prs, "Appendix A1 · Why tonnes, not percentiles",
        f"Percentile pillars squeeze a {ratio:.1f}x carbon gap into 10 points.",
        "Sources: engine pillar scores on main 23bb71c (docs/deck_build_log.md step 1); research/impact.md (dry cooling, "
        "eGRID2023 subregion rates); docs/weighting.md.",
        "A1",
        notes=f"""
If asked why we price the leaders in dollars and tonnes: at the county table's regional
grid rates, Franklin, New York emits {ratio:.1f} times less CO2 than Grant, but their energy
and carbon pillars score {f_score} and {g_score}, the {f_pct:.0f}th and {g_pct:.0f}th national
percentiles. The pillar blends grid rate with nearby clean capacity, where Grant scores near
the top, and percentiles compress large physical gaps. So tonnes and dollars need their own
model. One caveat: with BPA-like supply, Grant is about {GR['co2_tonnes']['bpa'] / 1000:.0f}
thousand tons, below Franklin.
""",
        checks=[("Energy and carbon pillar, Grant", f"{g_score} ({g_pct}th pctl)", "docs/deck_build_log.md", "step 1 (engine on main 23bb71c)", True),
                ("Energy and carbon pillar, Franklin", f"{f_score} ({f_pct}th pctl)", "docs/deck_build_log.md", "step 1 (engine on main 23bb71c)", True),
                ("CO2, Grant at Northwest average", f"{fmt(co2_g)} t/yr", "docs/figures/facts.json", "grant_ranges.co2_tonnes.nwpp_table", True),
                ("CO2, Franklin", f"{fmt(FRANKLIN_CO2)} t/yr", "research/impact.md", "Dry cooling table", True),
                ("Ratio", f"{ratio:.2f}x", "computed", None, True)],
    )
    half = int((CW - Inches(0.5)) / 2)
    bar_chart(s, MX, TOP, half, Inches(3.4), ["Franklin, NY", "Grant, WA"], [f_score, g_score], [BLUE, NAVY],
              title="Energy and carbon pillar score", fmt_code="0.0", max_val=100, cat_size=17)
    bar_chart(s, MX + half + Inches(0.5), TOP, half, Inches(3.4), ["Franklin, NY", "Grant, WA"],
              [FRANKLIN_CO2 / 1000, co2_g / 1000], [BLUE, NAVY],
              title="CO2, thousand metric tons a year", max_val=850, cat_size=17)
    bullets(s, MX, TOP + Inches(3.6), CW, Inches(1.1),
            ["The pillar blends grid rate with nearby clean capacity.",
             f"With BPA-like supply, Grant is {GR['co2_tonnes']['bpa'] / 1000:.0f}k t, below Franklin."], size=19, gap=8)


def app_tie(prs):
    s = frame(
        prs, "Appendix A2 · Three-county tie",
        "In 25-year dollars, three counties tie within 1.1%. Each has a win condition.",
        "Sources: docs/weighting.md (Recommendation; What each county needs to win; With and without tax), "
        "scratch/weighting/out/monetize_summary.json and montecarlo_summary*.json. 300 MW IT, 25 years at 7%, $190/t CO2.",
        "A2",
        notes="""
The dollar model prices energy, carbon at 190 dollars a ton, water, hazard, delay, and
sales tax over 25 years. Clark, Grant, and Franklin land within 1.1 percent, so the model
can't separate them. Grant wins if its power arrives within about 10 to 12 months of a
two-year baseline. Clark wins if its parcel avoids the transit district's higher tax.
Franklin wins if New York's data center tax exemption applies. With sales tax, full-exemption
states price cheaper nationally. Those flags are unverified, so checking them comes first.
""",
        checks=[(f"25-year cost, {k}", f"${v}B", "docs/weighting.md", "Recommendation table", True) for k, v in TIE.items()]
        + [("Grant break-even months vs Clark / Franklin", "10.2 / 12.1", "docs/weighting.md", "Grant's row", True),
           ("Shortlist ranks with tax: Clark, Grant, Franklin", "73, 78, 80 of 162", "docs/weighting.md", "The framework", True),
           ("Monte Carlo #1 with tax", "Chesterfield SC 21.8% of draws", "scratch/weighting/out/montecarlo_summary.json", "p_first_top", True)],
    )
    rows = [("County", "25-year cost", "Wins when"),
            ("Grant, WA (featured)", f"${TIE['Grant, WA']}B", "Full 300 MW within about 10 to 12 months of a 2-year baseline."),
            ("Clark, WA", f"${TIE['Clark, WA']}B", "Parcel outside the transit district (8.0% sales tax) and power in about 2.25 years."),
            ("Franklin, NY", f"${TIE['Franklin, NY']}B", "New York's data center sales tax exemption applies to the campus.")]
    table(s, MX, TOP, CW, Inches(2.9), rows, [Inches(3.0), Inches(2.2), CW - Inches(5.2)], size=17)
    bullets(s, MX, TOP + Inches(3.15), CW, Inches(1.6),
            ["With tax, full-exemption states price cheaper; those flags are unverified.",
             "Shortlist ranks with tax: Clark 73, Grant 78, Franklin 80."], size=19, gap=8)


def app_weights(prs):
    methods = ["Balanced", "Monetized $190/t", "CRITIC", "Entropy", "Revealed pref."]
    table_w = {  # docs/weighting.md, Weights by method
        "energy_carbon": [0.153, 0.228, 0.185, 0.249, 0.171], "water": [0.119, 0, 0.120, 0.038, 0.123],
        "climate_resilience": [0.119, 0.010, 0.252, 0.094, 0.224], "grid_infrastructure": [0.153, 0.066, 0.150, 0.441, 0.086],
        "land": [0.068, 0, 0.044, 0.008, 0.156], "community": [0.085, 0, 0.136, 0.144, 0.129],
        "permitting": [0.153, 0.155, 0.096, 0.024, 0.053], "cost": [0.150, 0.541, 0.018, 0.002, 0.058]}
    s = frame(
        prs, "Appendix A3 · Weights by method",
        "Five ways to weight the pillars disagree. That's why we test, not tune.",
        "Sources: docs/weighting.md (Weights by method); scratch/weighting/out/. Monetized: share of 25-year cost variance at "
        "$190/t with tax, negatives set to 0. CRITIC and entropy: column weights summed by pillar.",
        "A3",
        notes="""
We compared five ways to set weights. Balanced is our stated judgment. The monetized column
is how much each pillar's dollars vary across counties: cost of power dominates. CRITIC and
entropy derive weights from the data's spread, and entropy is an artifact here. Revealed
preference fits where industry already built. They disagree a lot, and no one of them is
right, which is why we report SMAA over every possible weighting instead of defending one.
""",
        checks=[("Weights table", "5 methods x 8 pillars", "docs/weighting.md", "Weights by method", True)],
    )
    cd = CategoryChartData()
    cd.categories = [lbl for _, lbl, _ in PILLARS]
    for i, m in enumerate(methods):
        cd.add_series(m, [table_w[k][i] * 100 for k, _, _ in PILLARS])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, MX, TOP, CW, Inches(4.7), cd)
    ch = gf.chart
    ch.font.name = FONT
    ch.font.size = Pt(15)
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(15)
    for ser, c in zip(ch.plots[0].series, [INK2, NAVY, GREEN, GREY, EMBER]):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = c
    ch.plots[0].gap_width = 60
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RULE
    va.tick_labels.font.size = Pt(14)
    va.tick_labels.number_format = '0"%"'
    va.tick_labels.number_format_is_linked = False
    va.format.line.fill.background()
    ch.category_axis.tick_labels.font.size = Pt(15)
    ch.category_axis.tick_labels.font.color.rgb = INK
    ch.category_axis.major_tick_mark = XL_TICK_MARK.NONE


def app_sources(prs):
    rows = [("Pillar", "Sources"),
            ("Energy and carbon", "eGRID2023 subregions; LBNL interconnection queue; NREL WIND Toolkit; eGRID plants within 100 km"),
            ("Water", "US Drought Monitor; FEMA NRI; WRI Aqueduct 4.0; CMRA"),
            ("Climate resilience", "FEMA NRI loss rates; CMRA projections (LOCA-downscaled CMIP5)"),
            ("Grid and infrastructure", "LBNL queue; FCC fiber; FracTracker; EIA-860"),
            ("Cost of power", "EIA-861 state industrial prices"),
            ("Land", "Census TIGER"),
            ("Community", "ACS; BLS LAUS; BEA 1969 employment; Census history since 1950"),
            ("Permitting", "EPA Green Book; hand-coded state policy tables")]
    s = frame(
        prs, "Appendix A4 · Data sources",
        f"Every county is scored from public data. Proxies are labeled as proxies.",
        "Sources: docs/deck.md (Data); data/processed/county_features.manifest.json (sources and versions); docs/schema.md; "
        "engine/pillars.yaml.",
        "A4",
        notes=f"""
All {fmt(B['counties'])} counties are scored from about 20 public sources, joined to one county
table. Four inputs are proxies, and we say so: last-mile fiber stands in for backbone, plant
capacity within 100 kilometers stands in for deliverable power, a state average price stands
in for what a new load pays, and a heat-sink score stands in for heat reuse. Coverage across
the top 10 is 73 to 78 percent, because 10 of 45 mapped columns aren't in the table yet.
""",
        checks=[("Counties", fmt(B["counties"]), "docs/figures/facts.json", "balanced.counties", True),
                ("Scored columns present / mapped", f"{F['limitations']['scored_columns_present']} / {F['limitations']['scored_columns_mapped']}",
                 "docs/figures/facts.json", "limitations.scored_columns_present, scored_columns_mapped", True),
                ("Coverage, top 10", f"{min(F['limitations']['coverage_top10']):.2f} to {max(F['limitations']['coverage_top10']):.2f}",
                 "docs/figures/facts.json", "limitations.coverage_top10", True)],
    )
    table(s, MX, TOP, CW, Inches(4.7), rows, [Inches(3.2), CW - Inches(3.2)], size=16)


def app_limits(prs):
    L = F["limitations"]
    s = frame(
        prs, "Appendix A5 · Limitations",
        "The engine screens on what's installed and average. Feasibility decides.",
        "Sources: docs/deck.md (Limitations); docs/figures/facts.json (limitations, weights); research/risk.md; "
        "research/implementation.md; docs/conditions.md.",
        "A5",
        notes=f"""
Three limits matter most. First, installed isn't available: Grant has {fmt(L['featured_plant_capacity_mw_100km'])}
megawatts of plants within 100 kilometers, yet its utility has no spare hydro for a new load.
Second, price is a state average, so the cost pillar ranks states, and a new 300 megawatt
load pays more. Third, some inputs are proxies, permitting is three hand-coded state-level
values, and the weights are judgments, which is why we test them. First place leads by
{F['gap_first_to_second']} points, so robustness carries the claim, not the rank.
""",
        checks=[("Plant capacity within 100 km of Grant", f"{fmt(L['featured_plant_capacity_mw_100km'])} MW", "docs/figures/facts.json", "limitations.featured_plant_capacity_mw_100km", True),
                ("Lead over #2", f"{F['gap_first_to_second']} points", "docs/figures/facts.json", "gap_first_to_second", True),
                ("Loudoun VA", L["loudoun_va"], "docs/figures/facts.json", "limitations.loudoun_va", True)],
    )
    bullets(s, MX, TOP + Inches(0.1), Inches(7.0), Inches(4.5),
            [f"Installed isn't available: {fmt(L['featured_plant_capacity_mw_100km'])} MW nearby, none spare.",
             "Price is a state average; a new large load pays more.",
             f"First place leads by {F['gap_first_to_second']} points; robustness carries the claim."],
            size=21, gap=18)
    rows = [("Proxy", "Stands in for"),
            ("Last-mile fiber", "Backbone fiber"),
            ("Plants within 100 km", "Deliverable power"),
            ("State average price", "New-load rate"),
            ("Heat-sink score", "Heat reuse"),
            ("Queue age", "Time to power"),
            ("State permitting values", "Local permitting risk")]
    table(s, MX + Inches(7.4), TOP + Inches(0.1), CW - Inches(7.4), Inches(4.4), rows,
          [Inches(2.2), CW - Inches(9.6)], size=14)


def app_extends(prs):
    s = frame(
        prs, "Appendix A6 · The engine extends",
        "Same engine, new region or new question: countries, and industrial reuse.",
        "Sources: docs/global.md; results/global_balanced.csv; docs/figures/global_table.png; docs/industrial_reuse.md; "
        "docs/demo_script.md (Boone County, IL; BLS QCEW private manufacturing jobs).",
        "A6",
        notes=f"""
Two extensions. First, the same engine ranked {GL['countries']} countries after about 40 lines of change:
{GL['passed']} pass the gates, and Sweden, Switzerland, and Norway lead. That shows the engine carries to
another region; it isn't a country recommendation, and there's no global power price. Second,
Stage 2 adds unscored context after the ranking. Boone County, Illinois lost almost three quarters of
its manufacturing jobs after the Belvidere plant went idle. Economic need is not evidence of
community support; local engagement is still required.
""",
        checks=[("Countries / pass / floor", f"{GL['countries']} / {GL['passed']} / {GL['floor_ok']}", "docs/figures/facts.json", "global", True),
                ("US global rank", f"{GL['us']['rank']} of {GL['us']['of']}", "docs/figures/facts.json", "global.us", True),
                ("Boone IL manufacturing jobs", "7,761 (2015) to 2,070 (2024)", "docs/demo_script.md", "Stage 2 table", True)],
    )
    half = int((CW - Inches(0.5)) / 2)
    with open(ROOT / "results/global_balanced.csv") as fh:
        glob = list(csv.DictReader(fh))
    us = next(r for r in glob if r["iso3"] == "USA")
    rows = [("Rank", "Country", "Score", "Floor")]
    for r in glob[:5] + [us]:
        rows.append((r["rank"], r["country"], f"{float(r['composite']):.1f}", "passes" if r["floor_ok"] == "True" else "fails"))
    textbox(s, MX, TOP - Inches(0.05), half, Inches(0.4),
            f"{GL['countries']} countries: {GL['passed']} pass the gates", size=18, bold=True)
    table(s, MX, TOP + Inches(0.45), half, Inches(3.3), rows,
          [Inches(1.0), half - Inches(3.4), Inches(1.2), Inches(1.2)], size=16)
    bar_chart(s, MX + half + Inches(0.5), TOP, half, Inches(3.4), ["2015", "2024"], [7761, 2070], [GREY, EMBER],
              title="Boone County, IL: manufacturing jobs", horizontal=False, max_val=9000, cat_size=17)
    textbox(s, MX + half + Inches(0.5), TOP + Inches(3.5), half, Inches(1.0),
            "Economic need is not evidence of community support.", size=18, bold=True, color=INK)


def table(slide, x, y, w, h, rows, col_w, size=16):
    gf = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, h)
    tbl = gf.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = cw
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (GROUND if i % 2 == 0 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = val
            style_run(r, size, WHITE if i == 0 else INK, bold=(i == 0 or j == 0))
    return tbl


# ---------------------------------------------------------------------------
# Companion files
# ---------------------------------------------------------------------------
def write_checks():
    lines = ["# Numbers to check", "",
             "Generated by `docs/deck/build_deck.py`. Every number below appears on a slide or in its",
             "speaker notes with `[[CHECK]]`, because it comes from the engine, the county table, or a",
             "model run that could change tonight. Recheck each against its source before presenting.",
             "Numbers from published external sources (statutes, the Grant PUD FAQ) are listed in the",
             "notes as stable and aren't repeated here.", "",
             "| Slide | Number | Value in the deck | Source file | Field |", "| --- | --- | --- | --- | --- |"]
    for slide, what, value, f, field in CHECKS:
        lines.append(f"| {slide} | {what} | {value} | `{f}` | {field or ''} |")
    (DECK / "numbers_to_check.md").write_text("\n".join(lines) + "\n")


def write_sources():
    lines = ["# Deck sources", "",
             "Generated by `docs/deck/build_deck.py`. Quotes, external sources, and image credits for",
             "`pitch_template.pptx`.", "",
             "## Quotes", "",
             "Every quote was checked on 2026-10-04 by opening the primary source and matching the words",
             "exactly, ignoring only whitespace and straight versus curly quotation marks. An ellipsis marks",
             "an omission that doesn't change the meaning. Quotes marked unused are verified and available",
             "for Q&A or a swap.", ""]
    for q in QUOTE_BANK:
        where = q["slide"].capitalize() if q["slide"].startswith("unused") else f"Slide {q['slide']}"
        lines += [f"- **{where}:** “{q['text']}”",
                  f"  - Attribution: {q['who']}",
                  f"  - Location: {q['source']}",
                  f"  - Link: {q['url']}"]
    lines += ["", "Not used, and why:", "",
              "- “at its water right limits” (City of Quincy): the words come from a Department of Ecology",
              "  meeting summary describing remarks by Bob Davis of the City of Quincy, not a City statement,",
              "  and the same summary attributes 60% of Quincy's water budget to food processing, not data",
              "  centers. Cite it as “Summary of City of Quincy remarks, CRPAG meeting, Oct. 23, 2025” if used.",
              "  https://www.ezview.wa.gov/Portals/_1962/Documents/CRPAG/Oct2025meetingum.pdf",
              "- No verified source sentence says siting choices matter “for decades”. The slide 1 headline",
              "  is the team's claim, not a quote.", "",
              "## Images", "",
              "- `docs/figures/map_composite.png` (slide 2): team figure from `docs/figures/make_figures.py`.",
              "- Every other chart is a native PowerPoint chart built from repo data, so its numbers can be edited.", "",
              "Openly licensed photos, checked on each Commons file page on 2026-10-04. None is placed on a",
              "slide yet; every main slide already carries a chart. Print the credit line in the slide's source",
              "line if you add one.", ""]
    for key, p in PHOTOS.items():
        lines.append(f"- {key}: {p['credit']}. License: {p['license']}. {p['page']}")
    lines += ["", "## External sources by slide", ""]
    for slide, srcs in SOURCES.items():
        for src in srcs:
            lines.append(f"- Slide {slide}: {src}")
    (DECK / "sources.md").write_text("\n".join(lines) + "\n")


if __name__ == "__main__":
    build()
